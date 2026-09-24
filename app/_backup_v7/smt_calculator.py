"""
smt_calculator.py
------------------
Deterministic implementation of the BMC Helix SMT's formula chains, so
the agent calls this code for arithmetic instead of reasoning through
multi-table lookups and discount math in prose.

Two entry points:
  - compute_new_business_line(...)  -> Part 1, New Input Form
  - compute_renewal_line(...)       -> Part 2, Renewal Input Form

Both raise ValueError with a clear message when a required input is
missing or a product can't be matched — the caller should surface that
as a flag to the user, never silently substitute a guess.
"""

SKILL_VERSION = 7

import csv
import os
from datetime import date, datetime

REF_DIR = os.path.join(os.path.dirname(__file__), "..", "references")


def _load_csv(filename):
    path = os.path.join(REF_DIR, filename)
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _normalize(name):
    return "".join(name.split()).lower()


class KnowledgeBase:
    """Loads all reference CSVs once; reused across line-item calcs."""

    def __init__(self):
        self.price_list = _load_csv("bmc_price_list.csv")
        self.channel_table = _load_csv("channel_discount_table.csv")
        self.sdn_table = _load_csv("sdn_disnet_gss_discount_table.csv")
        self.ervfactor_table = _load_csv("ervfactor_table.csv")
        self.trvfactor_table = _load_csv("trvfactor_table.csv")
        self.obsoleting = _load_csv("obsoleting_table.csv")
        self.cap_discount = _load_csv("cap_discount_table.csv")

        self._price_index = {}
        for row in self.price_list:
            key = (_normalize(row["Product Name"]), row["License Type"])
            self._price_index[key] = row

        self._obsoleting_index = {
            _normalize(r["Product Name"]): r["Projected Obsolete Date"]
            for r in self.obsoleting
        }
        self._cap_index = {
            _normalize(r["Product Name"]): float(r["Cap Discount %"])
            for r in self.cap_discount
        }

    def find_product(self, product_name, license_type):
        """Exact match required (after whitespace normalization).
        Returns None if not found — caller must flag, not guess."""
        key = (_normalize(product_name), license_type)
        return self._price_index.get(key)

    def channel_discount(self, product_family, customer_tier):
        for row in self.channel_table:
            if row["PF"] == product_family:
                val = row.get(customer_tier)
                if val not in (None, ""):
                    return float(val)
        return None

    def sdn_discount(self, product_line, deal_type, license_type):
        """deal_type: 'SDN' / 'Disnet' / 'GSS'. license_type: 'OPS' or 'Perpetual'."""
        col = f"{deal_type} % ({license_type})"
        for row in self.sdn_table:
            if row["Product Line (PL)"] == product_line:
                val = row.get(col)
                return float(val) if val not in (None, "") else 0.0
        return 0.0

    def obsolete_date(self, product_name):
        return self._obsoleting_index.get(_normalize(product_name))

    def cap_discount_pct(self, product_name):
        return self._cap_index.get(_normalize(product_name))

    def growth_factor(self, table, old_term_years, new_term_years):
        old_term_years = max(1, min(5, round(old_term_years)))
        new_term_years = max(1, min(5, round(new_term_years)))
        for row in table:
            if int(row["OldTermYrs"]) == old_term_years:
                return float(row[f"NewTermYrs_{new_term_years}"])
        raise ValueError(f"No growth factor row for OldTermYrs={old_term_years}")


# ---------------------------------------------------------------------
# PART 1 — New Input Form (genuinely new business)
# ---------------------------------------------------------------------

def compute_new_business_line(kb: KnowledgeBase, product_name, qty, license_type,
                                customer_tier, deal_type, term_months=12):
    """
    license_type: 'OPS' or 'Perpetual'
    deal_type: 'SDN', 'Disnet', 'GSS', or None. None raises rather than
               falling into the renewal branch, which requires Part 2
               data this function doesn't have.
    """
    row = kb.find_product(product_name, license_type)
    if row is None:
        raise ValueError(f"No exact price-list match for '{product_name}' ({license_type}). "
                          f"Flag this — do not guess a substitute.")

    list_price = float(row["List Price (USD)"])
    product_line = row["Product Line (PL)"]
    product_family = row["Product Family"]

    channel_disc = kb.channel_discount(product_family, customer_tier)
    if channel_disc is None:
        raise ValueError(f"No channel discount found for family='{product_family}', "
                          f"tier='{customer_tier}'. Flag — do not assume 0.")

    partner_price = list_price * (1 - channel_disc)

    if deal_type is None:
        raise ValueError("Deal Type (SDN/Disnet/GSS) is required and was not provided. "
                          "This is a per-deal input — ask, do not default.")

    bw = kb.sdn_discount(product_line, deal_type, license_type)

    if bw == 0:
        raise ValueError(
            f"BW=0 for Product Line '{product_line}' under Deal Type '{deal_type}' — "
            "flag this line for manual review rather than computing AS/AX automatically."
        )

    AS = 1 - (1 - bw) / (1 - channel_disc)
    AT = 1 - (1 - (0.9999 if deal_type == "GSS" else 0.9998)) / (1 - bw)

    AX = partner_price * (1 - AS) * (1 - AT)
    AZ = AX * term_months * qty
    BB = AZ
    AK = list_price * 12 * qty

    obs_date = kb.obsolete_date(product_name)

    return {
        "product": product_name, "license_type": license_type,
        "product_line": product_line, "product_family": product_family,
        "list_price": list_price, "partner_price": partner_price,
        "channel_discount": channel_disc, "BW": bw, "AS": AS, "AT": AT,
        "AX": AX, "AZ": AZ, "BA": 0.0, "BB": BB, "AK": AK,
        "obsolete_date": obs_date,
    }


# ---------------------------------------------------------------------
# PART 2 — Renewal Input Form (existing subscription)
# ---------------------------------------------------------------------

def compute_renewal_line(kb: KnowledgeBase, product_name, license_type,
                           existing_qty, proposed_qty, code,
                           orig_brv_acv, orig_trv_acv, orig_erv_acv, orig_term_months,
                           new_term_months, customer_tier,
                           support_rate=None,
                           obsolete_date=None, max_term_end_date_across_quote=None,
                           new_term_years_rounded_across_quote=None,
                           u_override=None, v_override=None, att_a_unit_cost=None):
    """
    code: 'a' (add), 'd' (decrease), 'x' (exclude), or '' (straight renewal)
    customer_tier: the exact channel_discount_table.csv column name, or
        the literal sentinel "Helix Base List Price", which disables
        channel discount entirely (0%).

        License Discount % (AS) and the channel/Customer Tier discount
        are two independent numbers that both multiply into AX. The
        channel discount is applied to list_price first, via `basis`;
        AS is applied to `basis` next. Both are returned explicitly.
    orig_brv_acv: prior period's actual booked ACV ($) for this line.
    orig_trv_acv / orig_erv_acv / orig_term_months: from the customer's
        own prior-period UFR/subscription record; cannot be derived
        from the current quote PDF. Raises if not supplied.
    support_rate: required for Perpetual lines (pass None for OPS/SaaS).
        Varies by Support Tier — resolve via get_support_rate() first.
        Raises if a Perpetual line is called without it.
    obsolete_date / max_term_end_date_across_quote / new_term_years_rounded_across_quote:
        feed the growth-factor capping condition. CG$11 and NewTL are
        whole-quote aggregates (max term-end date and rounded new-term-
        years across every line in the renewal), computed once by the
        caller and passed into every call. If omitted, the raw
        (uncapped) factor is used.
    """
    if orig_brv_acv is None or orig_trv_acv is None or orig_erv_acv is None or orig_term_months is None:
        raise ValueError(
            "Prior-period BRV/TRV/ERV ACV and term data are required for a "
            "renewal calculation and cannot come from the current quote PDF. "
            "Ask for the customer's UFR/prior-subscription export before proceeding."
        )
    if license_type != "OPS" and code not in ("d", "x") and support_rate is None:
        raise ValueError(
            f"support_rate is required for a Perpetual line ('{product_name}') "
            "— it is NOT a fixed 20%. Call get_support_rate(support_tier) "
            "first and pass its result here."
        )

    row = kb.find_product(product_name, license_type)
    if row is None:
        raise ValueError(f"No exact price-list match for '{product_name}' ({license_type}).")

    list_price = float(row["List Price (USD)"])
    product_family = row["Product Family"]

    # Real sheet stores prior-period ACVs as whole-dollar figures.
    orig_brv_acv = round(orig_brv_acv)
    orig_trv_acv = round(orig_trv_acv)
    orig_erv_acv = round(orig_erv_acv)

    # "Remediate On Prem" family convention: both orig_trv_acv and
    # orig_erv_acv are set to the rounded ERV value.
    if product_family == "Remediate On Prem":
        orig_trv_acv = orig_erv_acv

    common_meta = {
        "product": product_name, "license_type": license_type,
        "PL": row["Product Line (PL)"], "product_family": product_family,
        "uom": row.get("Unit of Measure", ""), "status": row.get("Status", ""),
        "list_price": list_price, "new_term_months": new_term_months,
        "orig_brv_acv": orig_brv_acv, "orig_trv_acv": orig_trv_acv,
        "orig_erv_acv": orig_erv_acv, "orig_term_months": orig_term_months,
        "support_rate": support_rate,
    }

    if code in ("d", "x"):
        return {
            **common_meta,
            "code": code, "existing_qty": existing_qty, "proposed_qty": 0,
            "channel_discount": None, "basis": None, "effective_discount": 0.0,
            "AS": 0.0, "AT": 0.0, "AU": 0.0, "AX": list_price, "AY": 0.0,
            "AZ": 0.0, "BA": 0.0, "BB": 0.0, "ERVFactor": None, "TRVFactor": None,
            "AQ": None,
            "note": "Code is decrease/exclude — real sheet zeroes AZ/BA/BB. "
                    "If this quantity reappears as a new OPS/Perpetual line "
                    "elsewhere in the same quote, flag as a possible license "
                    "conversion, not a genuine reduction.",
        }

    # Growth factors (old-term-years -> new-term-years lookup)
    old_yrs = orig_term_months / 12
    new_yrs = new_term_months / 12
    erv_factor = kb.growth_factor(kb.ervfactor_table, old_yrs, new_yrs)
    trv_factor_raw = kb.growth_factor(kb.trvfactor_table, old_yrs, new_yrs)

    # AO blank (no obsolescence entry) behaves as "greater than any
    # date" in Excel, so the AND is FALSE and the raw factor applies
    # whenever obsolete_date is None.
    cap_condition = (
        obsolete_date is not None
        and max_term_end_date_across_quote is not None
        and new_term_years_rounded_across_quote is not None
        and obsolete_date < max_term_end_date_across_quote
        and round(new_yrs) < new_term_years_rounded_across_quote
    )
    trv_factor = min(trv_factor_raw, 1) if cap_condition else trv_factor_raw

    bn = list_price * 12 if license_type == "OPS" else list_price * support_rate
    bo = max(filter(None, [u_override, v_override])) if (u_override or v_override) else bn

    brv_unit = orig_brv_acv / existing_qty if existing_qty else 0.0
    trv_unit = orig_trv_acv / existing_qty if existing_qty else 0.0

    # Y (Anchor Price TRV, per-unit) = MIN( brv_unit + (trv_unit - brv_unit) * factor, BN, BO )
    y = min(brv_unit + (trv_unit - brv_unit) * trv_factor, bn, bo)
    ad = y * proposed_qty
    ae = ad

    at = 0.0
    denom = (list_price * 12 * proposed_qty) if license_type == "OPS" else (list_price * support_rate * proposed_qty)
    AS = 1 - (ae / denom) / (1 - at) if denom else 0.0

    # channel_discount / basis ("Partner Discounted Price", column J/L):
    # the Customer Tier discount, applied to list_price before AS.
    channel_disc = None
    basis = list_price
    if customer_tier != "Helix Base List Price":
        row_for_family = row
        channel_disc = kb.channel_discount(row_for_family["Product Family"], customer_tier)
        if channel_disc is not None:
            basis = list_price * (1 - channel_disc)
    AX = basis * (1 - AS) * (1 - at)

    # Effective Discount % — combined Customer Tier + License Discount,
    # computed multiplicatively and derived straight from AX/list_price
    # so it can never drift out of sync with AX itself.
    effective_discount = 1 - (AX / list_price) if list_price else 0.0

    if license_type == "OPS":
        AZ = AX * new_term_months * proposed_qty
        BA = 0.0
        AQ = None
    else:
        AZ = AX * max(proposed_qty - existing_qty, 0)
        AQ = list_price * support_rate / 12
        AV = AS
        AW = 0.0
        AY = AQ * (1 - AV) * (1 - AW)
        BA = AY * proposed_qty * new_term_months

    BB = AZ + BA

    return {
        **common_meta,
        "code": code, "existing_qty": existing_qty, "proposed_qty": proposed_qty,
        "ERVFactor": erv_factor, "TRVFactor": trv_factor, "AQ": AQ,
        "channel_discount": channel_disc, "basis": basis,
        "effective_discount": effective_discount,
        "AS": AS, "AT": at, "AU": None, "AX": AX,
        "AY": (AY if license_type != "OPS" else 0.0),
        "AZ": AZ, "BA": BA, "BB": BB,
    }


# ---------------------------------------------------------------------
# xlsx OUTPUT — exact Renewal Input Form column layout
# ---------------------------------------------------------------------

# Exact header row, order, and labels from the real 'Renewal Input
# Form' tab. Do not reorder, rename, or drop columns — must stay
# directly paste-compatible into the real sheet.
RENEWAL_HEADERS = {
    "A": "Product (Marketing Schedule)", "B": "Product Number", "C": "Charge Number",
    "D": "Subscription Number", "E": "Status", "F": "PL", "G": "Product Family",
    "H": "UoM", "I": "BMC Base List Price ", "J": "Partner Discounted Price",
    "K": "BMC Base List Price ", "L": "Partner Discounted Price",
    "M": "Perp or OPS   (if applicable)",
    "N": "Approximate Rate Plan Name \n(ensure correct Revenue/Service Type tags if copying)",
    "O": "PRP Number (ensure correct Serv Type <--)", "P": "Code",
    "Q": "Existing\nQty", "R": "Proposed\nQty",
    "S": "Installmt/ Ramp? (edit if needed)",
    "U": '"Att A Unit Cost" (if applicable)', "V": "TRV Override (e.g., Price Capped)",
    "W": "BRV", "X": "'Hold the Line'\nERV", "Y": "'Anchor Price'\nTRV",
    "Z": "Proposed (incl all discounts)", "AA": "Calculated",
    "AB": "BRV (ACV)", "AC": "'Hold the Line'\nERV (ACV)",
    "AD": "'Anchor Price'\nTRV (ACV)", "AE": "Proposed (incl all discounts)",
    "AF": "Calculated", "AG": "On Anchor", "AH": "Calculated",
    "AI": "Estimated ERV Deviation", "AJ": "Estimated TRV Deviation",
    "AK": '(BP Use)      "Total Annual List"', "AL": "Start Date", "AM": "End Date",
    "AN": "Approx. Term in Months", "AO": "Projected Obsolete Date",
    "AP": "Perp Support Rate", "AQ": "Est. Perp Supp Mthly List Price",
    "AR": "Hold Disc", "AS": "License\nTrans", "AT": "Lic One-Time",
    "AU": "Total License", "AV": "Supp Trans", "AW": "Supp One-Time",
    "AX": "License Net Price", "AY": "Support Net Price", "AZ": "Total License",
    "BA": "Total Support", "BB": "Total", "BC": "orig TRV ACV",
    "BD": "orig ERV ACV", "BE": "orig term length", "BF": "chosen unit list",
    "BG": "code", "BH": "sum orig TRV ACV", "BI": "sum exist qty",
    "BJ": "sum decr qty", "BK": "sum rem qty", "BL": "sum new qty",
    "BM": "orig TRV ACV unit", "BN": "chosen annual unit list",
    "BO": "override unit if any", "BP": "new TRV ACV unit", "BQ": "decrease count",
    "BR": "disc form w vlookup", "BS": "max disc", "BT": "total trans disc",
    "BU": "disc cap", "BV": "total disc excl ARG", "BW": "orig end date",
    "BX": "dedup ramp/install qty count", "BY": "dedup ramp/install qty",
    "BZ": "Obsolete During Period", "CA": "SaaS margin", "CB": "SaaS location",
    "CC": "SaaS cost code", "CD": "SaaS margin cost", "CE": "deal score type",
}
RENEWAL_LAST_COL = "CE"

# Appended after CE as a reference-only audit trail for the Customer
# Tier discount and combined Effective Discount %, not part of the
# real template. Given a distinct header fill; strip before pasting
# into the real sheet.
REFERENCE_HEADERS = {
    "CF": "Customer Tier Discount %\n(reference only — not in SMT template)",
    "CG": "Effective Discount %\n(reference only — Tier + License combined)",
}
REFERENCE_LAST_COL = "CG"


def write_renewal_output_xlsx(results, output_path, start_date=None, end_date=None,
                                subscription_number=None):
    """
    results: list of dicts, each the return value of compute_renewal_line(),
        in the order they should appear as rows.
    output_path: where to save the .xlsx.
    start_date / end_date: the quote's contract dates, applied to every
        row's Start Date / End Date columns.
    subscription_number: applied to every row's Subscription Number
        column if provided.

    Populates only the columns this calculator computes or was given
    as input; every other column is left blank. CF/CG are appended
    after the real template's last column (CE) as a reference-only
    audit trail — see REFERENCE_HEADERS.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import column_index_from_string, get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Renewal Input Form"

    FONT = "Arial"
    HEADER_FILL = PatternFill("solid", fgColor="0064FF")
    HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=8)
    REF_HEADER_FILL = PatternFill("solid", fgColor="7D7D7D")
    DATA_FONT = Font(name=FONT, size=9)
    NA_FONT = Font(name=FONT, italic=True, size=9, color="808080")
    thin = Side(style="thin", color="D9D9D9")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

    last_idx = column_index_from_string(RENEWAL_LAST_COL)
    ref_last_idx = column_index_from_string(REFERENCE_LAST_COL)
    hdr_row = 1
    for c in range(1, ref_last_idx + 1):
        col = get_column_letter(c)
        is_reference_col = c > last_idx
        cell = ws.cell(
            row=hdr_row, column=c,
            value=RENEWAL_HEADERS.get(col) or REFERENCE_HEADERS.get(col, ""),
        )
        cell.fill = REF_HEADER_FILL if is_reference_col else HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[hdr_row].height = 42

    def setcell(row, col, value, font=DATA_FONT, fmt=None):
        c = ws.cell(row=row, column=column_index_from_string(col), value=value)
        c.font = font
        if fmt:
            c.number_format = fmt
        c.border = BORDER
        return c

    start_row = hdr_row + 1
    for idx, r in enumerate(results):
        row = start_row + idx
        license_type = r.get("license_type", "")
        is_ops = license_type == "OPS"

        setcell(row, "A", r.get("product", ""))
        if subscription_number:
            setcell(row, "D", subscription_number)
        setcell(row, "E", r.get("status", ""))
        setcell(row, "F", r.get("PL", ""))
        setcell(row, "G", r.get("product_family", ""))
        setcell(row, "H", r.get("uom", ""))

        list_price = r.get("list_price")
        basis = r.get("basis")
        if license_type == "Perpetual":
            setcell(row, "I", list_price, fmt='$#,##0.00000')
            if basis is not None:
                setcell(row, "J", basis, fmt='$#,##0.00000')
        elif is_ops:
            setcell(row, "K", list_price, fmt='$#,##0.00000')
            if basis is not None:
                setcell(row, "L", basis, fmt='$#,##0.00000')

        setcell(row, "M", license_type)
        setcell(row, "P", r.get("code", ""))
        setcell(row, "Q", r.get("existing_qty"))
        setcell(row, "R", r.get("proposed_qty"))

        if r.get("orig_brv_acv") is not None:
            setcell(row, "AB", r["orig_brv_acv"], fmt='$#,##0')

        setcell(row, "BC", r.get("orig_trv_acv"), fmt='$#,##0')
        setcell(row, "BD", r.get("orig_erv_acv"), fmt='$#,##0')
        setcell(row, "BE", r.get("orig_term_months"), fmt='0.0000')

        if start_date:
            setcell(row, "AL", start_date, fmt='DD-MMM-YYYY')
        if end_date:
            setcell(row, "AM", end_date, fmt='DD-MMM-YYYY')
        setcell(row, "AN", r.get("new_term_months"), fmt='0.0000')

        if license_type == "Perpetual":
            setcell(row, "AP", r.get("support_rate") or 0.0, fmt='0%')
            if r.get("AQ") is not None:
                setcell(row, "AQ", r["AQ"], fmt='$#,##0.00000')

        if r.get("AS") is not None:
            setcell(row, "AS", r["AS"], fmt='0.0000%')
        if r.get("AT") is not None:
            setcell(row, "AT", r["AT"], fmt='0.0000%')
        if r.get("AX") is not None:
            setcell(row, "AX", r["AX"], fmt='$#,##0.000000')
        if r.get("AY") is not None:
            setcell(row, "AY", r["AY"], fmt='$#,##0.000000')
        setcell(row, "AZ", r.get("AZ"), Font(name=FONT, bold=True, size=9), fmt='$#,##0.00')
        setcell(row, "BA", r.get("BA"), Font(name=FONT, bold=True, size=9), fmt='$#,##0.00')
        setcell(row, "BB", r.get("BB"), Font(name=FONT, bold=True, size=9), fmt='$#,##0.00')

        if r.get("note"):
            setcell(row, "BF", r["note"], NA_FONT)

        if r.get("channel_discount") is not None:
            setcell(row, "CF", r["channel_discount"], fmt='0.0000%')
        if r.get("effective_discount") is not None:
            setcell(row, "CG", r["effective_discount"], Font(name=FONT, bold=True, size=9), fmt='0.0000%')

        for c in range(1, ref_last_idx + 1):
            ws.cell(row=row, column=c).border = BORDER

    total_row = start_row + len(results)
    ws.cell(row=total_row, column=2, value="TOTAL").font = Font(bold=True)
    bb_col = column_index_from_string("BB")
    tot = ws.cell(row=total_row, column=bb_col,
                   value=f"=SUM({get_column_letter(bb_col)}{start_row}:{get_column_letter(bb_col)}{total_row-1})")
    tot.font = Font(bold=True)
    tot.number_format = '$#,##0.00'

    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 20
    ws.column_dimensions["BF"].width = 50
    ws.column_dimensions["CF"].width = 16
    ws.column_dimensions["CG"].width = 16
    for c in range(1, ref_last_idx + 1):
        col = get_column_letter(c)
        if col not in ("A", "G", "H", "BF", "CF", "CG"):
            ws.column_dimensions[col].width = 11
    ws.freeze_panes = f"A{start_row}"

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    # Smoke test: confirms both code paths run without errors. Does not
    # verify correctness.
    print(f"smt_calculator.py SKILL_VERSION={SKILL_VERSION}")
    print()
    kb = KnowledgeBase()

    print("=== Part 1 smoke test: BMC Discovery, SDN, Reseller Tier 1-3 ===")
    r = compute_new_business_line(kb, "BMC Discovery - Resource Unit", qty=1,
                                    license_type="OPS", customer_tier="Reseller (Customer Tier 1-3)",
                                    deal_type="SDN")
    print(f"AX={r['AX']:.6f}  BB={r['BB']:.4f}")

    print()
    print("=== Part 2 smoke test: Named User, straight renewal ===")
    r2 = compute_renewal_line(kb, "BMC Helix Service Management OnPrem - Standard - Named User",
                                license_type="OPS", existing_qty=50, proposed_qty=50, code="",
                                orig_brv_acv=807, orig_trv_acv=920, orig_erv_acv=892, orig_term_months=11.9667,
                                new_term_months=11, customer_tier="Helix Base List Price")
    print(f"AS={r2['AS']:.10f}")
    print(f"AX={r2['AX']:.4f}  BB={r2['BB']:.2f}")
    print(f"channel_discount={r2['channel_discount']}  basis={r2['basis']}")
    print(f"effective_discount={r2['effective_discount']:.10f}")
